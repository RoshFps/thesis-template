import datetime as dt
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
D = dt.date

FONT = "Arial"
f = lambda **k: Font(name=FONT, **k)
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

# (id, title, description, asset, category, owner, L, I, existing controls,
#  annex A controls, treatment, planned action, status, due, res L, res I)
RISKS = [
 ("R01", "Ransomware on central IT infrastructure",
  "Criminal group encrypts file servers, directory service and virtualisation hosts after initial access via phishing or an unpatched edge device. Comparable attacks hit several German universities between 2019 and 2023.",
  "Central servers, directory service, backups", "Cyber attack", "Head of University IT Centre (HIZ)", 4, 5,
  "Endpoint antivirus; nightly backups; firewall at network edge",
  "A.8.7 Protection against malware; A.8.13 Information backup; A.5.29 Information security during disruption",
  "Mitigate", "Introduce offline/immutable backup copy, tested quarterly restore, EDR on servers", "In progress", D(2026,12,31), 2, 4),
 ("R02", "Phishing and theft of staff or student credentials",
  "Credential phishing against university email accounts leads to mailbox takeover, internal spam waves and access to connected services via single sign-on.",
  "Email, single sign-on accounts", "Cyber attack", "Chief Information Security Officer (CISO)", 5, 3,
  "Spam filter; occasional awareness emails",
  "A.6.3 Information security awareness, education and training; A.8.5 Secure authentication; A.5.17 Authentication information",
  "Mitigate", "Roll out MFA for all staff accounts; yearly phishing simulation and e-learning", "In progress", D(2026,9,30), 3, 2),
 ("R03", "Leak of sensitive research data",
  "Personal or health data from research projects (for example clinical studies at the Homburg campus) is exposed through insecure storage, email or unauthorised sharing.",
  "Research data stores, project drives", "Data protection", "Vice President for Research", 3, 5,
  "Data protection training for researchers; ethics committee review",
  "A.5.12 Classification of information; A.5.34 Privacy and protection of PII; A.8.11 Data masking",
  "Mitigate", "Classification scheme for research data; encrypted research storage offering; pseudonymisation guideline", "Planned", D(2027,3,31), 2, 4),
 ("R04", "Loss or theft of laptops and mobile devices",
  "Staff laptops and phones holding exam papers, HR records or research data are lost while travelling or commuting; many devices are not centrally managed or encrypted.",
  "Endpoints", "Physical / human", "Chief Information Security Officer (CISO)", 4, 3,
  "Encryption recommended in IT guideline, not enforced",
  "A.8.1 User endpoint devices; A.7.9 Security of assets off-premises; A.8.24 Use of cryptography",
  "Mitigate", "Mandatory full-disk encryption and device management for university-owned laptops", "In progress", D(2026,8,31), 2, 2),
 ("R05", "Outage of a critical IT or network supplier",
  "Loss of the external network uplink (assumed to run via the German research network DFN, as is typical for German universities) or of a cloud collaboration service stops email, teaching platforms and remote access.",
  "Internet uplink, cloud services", "Third party", "Head of University IT Centre (HIZ)", 2, 4,
  "Contracts with standard SLAs; single uplink path",
  "A.5.19 Information security in supplier relationships; A.5.21 Managing security in the ICT supply chain; A.5.30 ICT readiness for business continuity",
  "Mitigate", "Supplier risk review for top 10 IT suppliers; document fallback for loss of uplink", "Open", D(2027,1,31), 2, 3),
 ("R06", "Unpatched servers run by institutes (shadow IT)",
  "Chairs and institutes operate their own servers and web applications outside central IT; no inventory exists and patches are applied irregularly.",
  "Decentral servers and web apps", "Vulnerability", "Chief Information Security Officer (CISO)", 4, 4,
  "Annual vulnerability scan of public IP ranges",
  "A.5.9 Inventory of information and other associated assets; A.8.8 Management of technical vulnerabilities",
  "Mitigate", "Asset inventory for all servers on the campus network; monthly external scanning with follow-up", "Open", D(2026,7,31), 2, 3),
 ("R07", "Campus management system failure during enrolment or exam registration",
  "The student and exam administration system becomes unavailable in a peak period, so students cannot enrol or register for exams before deadlines.",
  "Campus management system", "Availability", "Head of Student Services", 2, 4,
  "Vendor support contract; daily database backup",
  "A.8.14 Redundancy of information processing facilities; A.5.30 ICT readiness for business continuity; A.8.6 Capacity management",
  "Mitigate", "Load test before semester start; documented manual fallback and deadline extension procedure", "Planned", D(2027,2,28), 1, 3),
 ("R08", "Misuse of privileged administrator accounts",
  "An administrator, or an attacker who has taken over an admin account, abuses broad standing privileges to read data or disable controls without being noticed.",
  "Admin accounts, directory service", "Insider / access", "Head of University IT Centre (HIZ)", 2, 5,
  "Separate admin accounts for some systems",
  "A.8.2 Privileged access rights; A.8.15 Logging; A.5.18 Access rights",
  "Mitigate", "Privileged access review twice a year; central logging of admin actions", "Open", D(2027,3,31), 1, 4),
 ("R09", "Accounts not removed after people leave",
  "Accounts of former staff, guest researchers and exmatriculated students stay active because HR and student systems are not linked to identity management.",
  "Identity management", "Insider / access", "Director of Human Resources", 4, 3,
  "Manual deactivation on request",
  "A.5.16 Identity management; A.5.18 Access rights",
  "Mitigate", "Automated joiner-mover-leaver process fed by HR and student records", "In progress", D(2026,11,30), 2, 2),
 ("R10", "Personal data exposed through misconfigured file sharing",
  "Cloud folders or file shares with student or applicant data are shared via public links or with too broad groups.",
  "Cloud storage, file shares", "Data protection", "Data Protection Officer", 3, 4,
  "Data protection guideline; sharing links unrestricted by default",
  "A.5.14 Information transfer; A.5.15 Access control; A.5.34 Privacy and protection of PII",
  "Mitigate", "Default expiry on sharing links; quarterly report of public shares to owners", "Planned", D(2027,1,31), 2, 3),
 ("R11", "Espionage targeting research and knowledge transfer",
  "State-sponsored actors target research groups in computer science and engineering to steal unpublished results or export-controlled know-how.",
  "Research IP, lab networks", "Cyber attack", "Vice President for Research", 3, 4,
  "Export control officer; general firewall rules",
  "A.5.7 Threat intelligence; A.5.12 Classification of information; A.8.16 Monitoring activities",
  "Mitigate", "Security briefing for high-risk research groups; segmented network for sensitive projects", "Open", D(2027,6,30), 2, 3),
 ("R12", "Fire, water or power failure in the data centre",
  "A physical event in the central server room destroys or shuts down core systems; backups are kept in the same building.",
  "Data centre", "Physical / environmental", "Head of Facilities Management", 2, 5,
  "Fire detection; UPS for short outages",
  "A.7.5 Protecting against physical and environmental threats; A.7.11 Supporting utilities; A.8.13 Information backup",
  "Mitigate", "Move backup copy to a second site on campus; test UPS and generator yearly", "In progress", D(2026,10,31), 1, 4),
 ("R13", "Denial-of-service attack on website and network",
  "Overload attack on the public website, VPN or exam platform during a sensitive period such as online exams.",
  "Website, VPN, exam platform", "Cyber attack", "Head of University IT Centre (HIZ)", 3, 3,
  "Basic DDoS filtering by network provider",
  "A.8.20 Networks security; A.8.6 Capacity management",
  "Accept", "Accept residual risk; keep provider DDoS filter and document escalation contact", "Accepted", D(2026,6,30), 3, 3),
 ("R14", "Insecure software developed in-house",
  "Web portals and research tools built by institutes and student assistants contain common vulnerabilities (injection, broken access control) because no secure development rules exist.",
  "In-house web applications", "Vulnerability", "Chief Information Security Officer (CISO)", 3, 3,
  "None formalised",
  "A.8.25 Secure development life cycle; A.8.28 Secure coding; A.8.29 Security testing in development and acceptance",
  "Mitigate", "Short secure coding guideline; security check before a web app goes public", "Planned", D(2027,4,30), 2, 2),
 ("R15", "Attacks detected too late",
  "No central log collection or incident process means a compromise can stay undetected for weeks and response is improvised.",
  "Security monitoring", "Detection / response", "Chief Information Security Officer (CISO)", 4, 4,
  "Logs kept locally on some systems",
  "A.8.16 Monitoring activities; A.5.24 Information security incident management planning and preparation; A.5.26 Response to information security incidents",
  "Mitigate", "Central log platform for core systems; written incident response plan with on-call list", "In progress", D(2026,9,15), 2, 3),
 ("R16", "Insecure lab equipment and building systems on the campus network",
  "Lab devices, building automation and other connected equipment run outdated software and share the network with office computers.",
  "Lab and building IoT", "Vulnerability", "Head of Facilities Management", 3, 4,
  "Some VLAN separation",
  "A.8.22 Segregation of networks; A.8.20 Networks security",
  "Mitigate", "Separate network zones for building systems and lab devices", "Open", D(2027,6,30), 2, 3),
 ("R17", "Loss of key IT staff and knowledge",
  "Critical systems depend on a few long-serving administrators; documentation is thin and public-sector salaries make replacements hard to hire.",
  "IT operations", "People", "Head of University IT Centre (HIZ)", 3, 3,
  "Informal deputies for some systems",
  "A.5.37 Documented operating procedures; A.5.2 Information security roles and responsibilities",
  "Mitigate", "Runbooks for top 15 services; named deputy for every critical system", "Planned", D(2027,3,31), 2, 2),
 ("R18", "Confidential data entered into public AI tools",
  "Staff and students paste exam questions, personal data or unpublished research into public generative AI services without a contract or data protection review.",
  "Generative AI use", "Data protection", "Data Protection Officer", 4, 3,
  "No specific rule",
  "A.5.10 Acceptable use of information and other associated assets; A.5.14 Information transfer",
  "Mitigate", "AI usage guideline; offer a contracted AI service with data protection agreement", "In progress", D(2026,12,31), 2, 2),
]

wb = Workbook()

# ---------------- Read Me ----------------
rm = wb.active
rm.title = "Read Me"
rm.column_dimensions["A"].width = 110
lines = [
 ("Information Security Risk Register: Saarland University", f(bold=True, size=14, color="1F3864")),
 ("PORTFOLIO EXERCISE. Not an official document of Saarland University.", f(bold=True, color="C00000")),
 ("", None),
 ("Purpose", f(bold=True)),
 ("Sample risk register for an ISMS based on ISO/IEC 27001:2022, built around Saarland University (Universität des Saarlandes) as the example organisation.", f()),
 ("Public facts used: campuses in Saarbrücken and Homburg, the Faculty of Medicine in Homburg, and the University IT Centre (HIZ), shared with htw saar. A network uplink via the German research network (DFN) is assumed, not confirmed. Links are on the Sources sheet.", f()),
 ("Everything else (risks, scores, owners, existing controls, due dates, statuses) is invented for illustration and does not describe the university's real security posture.", f()),
 ("", None),
 ("Sheets", f(bold=True)),
 ("Risk Register: one row per risk with inherent and residual scoring, owner role, ISO 27001:2022 Annex A controls, treatment and remediation tracking.", f()),
 ("Summary: KPIs and a 5x5 heatmap, calculated from the register with formulas.", f()),
 ("Scoring Method: definitions for likelihood and impact (1 to 5), rating bands and risk appetite.", f()),
 ("Lists: values used by the drop-down menus.", f()),
 ("Sources: public pages and standards behind the facts, scales and control references.", f()),
 ("", None),
 ("How to edit", f(bold=True)),
 ("Cells with a light yellow fill are inputs. Grey-headed columns (Score, Rating, Days overdue, Overdue, Residual score) are formulas and should not be overwritten.", f()),
 ("Likelihood, Impact, Treatment and Status use drop-downs. Rating band thresholds live on the Scoring Method sheet and drive all ratings.", f()),
 ("Days overdue uses TODAY(), so it updates whenever the file is opened. Closed and Accepted risks are never counted as overdue.", f()),
 ("To add a risk, copy the last row down and extend the formulas; the Summary sheet counts rows 5 to 100.", f()),
 ("", None),
 (f"Version 0.1, prepared {dt.date(2026,10,8):%d %B %Y}.", f(italic=True, color="595959")),
]
for i, (text, font) in enumerate(lines, 1):
    c = rm.cell(row=i, column=1, value=text)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if font:
        c.font = font

# ---------------- Scoring Method ----------------
sm = wb.create_sheet("Scoring Method")
sm["A1"] = "Scoring Method"; sm["A1"].font = f(bold=True, size=14, color="1F3864")
sm["A2"] = "Score = Likelihood x Impact (1 to 25). Assumed scales for this portfolio exercise."
sm["A2"].font = f(italic=True)
hdr = ["Level", "Likelihood", "Meaning (likelihood)", "Impact", "Meaning (impact)"]
scale = [
 (1, "Rare", "Less than once in 10 years", "Negligible", "No noticeable effect on teaching, research or administration"),
 (2, "Unlikely", "Once in 5 to 10 years", "Minor", "Short disruption of one service; no personal data affected"),
 (3, "Possible", "Once in 2 to 5 years", "Moderate", "Disruption of several services for up to a day, or limited personal data breach"),
 (4, "Likely", "About once a year", "Major", "Core services down for days, reportable GDPR breach, or loss of research results"),
 (5, "Almost certain", "Several times a year", "Severe", "University-wide outage for weeks, large data breach, or serious reputational and legal damage"),
]
for j, h in enumerate(hdr, 1):
    c = sm.cell(row=4, column=j, value=h); c.font = f(bold=True, color="FFFFFF"); c.fill = HEAD_FILL; c.border = BORDER
for i, row in enumerate(scale, 5):
    for j, v in enumerate(row, 1):
        c = sm.cell(row=i, column=j, value=v); c.font = f(); c.border = BORDER; c.alignment = WRAP

sm["A11"] = "Rating bands (edit the yellow thresholds)"; sm["A11"].font = f(bold=True)
bands = [("Low", 1, "Accept; review yearly"),
         ("Medium", 5, "Treat when cost-effective; review every 6 months"),
         ("High", 10, "Treatment plan required; owner reports quarterly"),
         ("Critical", 15, "Immediate treatment; report to university management")]
for j, h in enumerate(["Rating", "Minimum score", "Required response"], 1):
    c = sm.cell(row=12, column=j, value=h); c.font = f(bold=True, color="FFFFFF"); c.fill = HEAD_FILL; c.border = BORDER
for i, (name, mn, resp) in enumerate(bands, 13):
    sm.cell(row=i, column=1, value=name)
    sm.cell(row=i, column=2, value=mn).fill = INPUT_FILL
    sm.cell(row=i, column=2).font = f(color="0000FF")
    sm.cell(row=i, column=3, value=resp)
    for j in (1, 3):
        sm.cell(row=i, column=j).font = f()
    for j in (1, 2, 3):
        sm.cell(row=i, column=j).border = BORDER
sm["A18"] = "Risk appetite: residual score must be below the High threshold (B15). Risks at or above it need a treatment plan or a documented acceptance by the risk owner."
sm["A18"].font = f(italic=True)
for col, w in zip("ABCDE", (16, 18, 34, 14, 70)):
    sm.column_dimensions[col].width = w

# ---------------- Lists ----------------
ls = wb.create_sheet("Lists")
lists = {
 "A": ("Score", [1, 2, 3, 4, 5]),
 "B": ("Treatment", ["Mitigate", "Transfer", "Avoid", "Accept"]),
 "C": ("Status", ["Open", "Planned", "In progress", "Closed", "Accepted"]),
}
for col, (h, vals) in lists.items():
    ls[f"{col}1"] = h; ls[f"{col}1"].font = f(bold=True)
    for i, v in enumerate(vals, 2):
        ls[f"{col}{i}"] = v; ls[f"{col}{i}"].font = f()
    ls.column_dimensions[col].width = 16

# ---------------- Risk Register ----------------
ws = wb.create_sheet("Risk Register", 1)
ws["A1"] = "Information Security Risk Register: Saarland University"
ws["A1"].font = f(bold=True, size=14, color="1F3864")
ws["A2"] = "Portfolio exercise based on ISO/IEC 27001:2022. Public facts only; all internal details are invented. Not an official university document."
ws["A2"].font = f(italic=True, color="C00000")

cols = [
 ("Risk ID", 8, "in"), ("Risk title", 30, "in"), ("Description (threat and vulnerability)", 55, "in"),
 ("Affected asset / process", 24, "in"), ("Category", 16, "in"), ("Risk owner (role)", 26, "in"),
 ("Likelihood (1-5)", 11, "in"), ("Impact (1-5)", 10, "in"), ("Inherent score", 10, "calc"), ("Inherent rating", 11, "calc"),
 ("Existing controls", 32, "in"), ("ISO 27001:2022 Annex A controls", 42, "in"), ("Treatment", 12, "in"),
 ("Planned remediation", 42, "in"), ("Status", 12, "in"), ("Due date", 12, "in"),
 ("Days overdue", 10, "calc"), ("Overdue", 9, "calc"),
 ("Residual likelihood", 11, "in"), ("Residual impact", 10, "in"), ("Residual score", 10, "calc"), ("Residual rating", 11, "calc"),
 ("Last reviewed", 12, "in"),
]
HR = 4
CALC_FILL = PatternFill("solid", fgColor="595959")
for j, (name, w, kind) in enumerate(cols, 1):
    c = ws.cell(row=HR, column=j, value=name)
    c.font = f(bold=True, color="FFFFFF")
    c.fill = CALC_FILL if kind == "calc" else HEAD_FILL
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    c.border = BORDER
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[HR].height = 45

def rating(score_cell):
    s = "'Scoring Method'!$B$"
    return (f'=IF({score_cell}="","",IF({score_cell}>={s}16,"Critical",IF({score_cell}>={s}15,"High",'
            f'IF({score_cell}>={s}14,"Medium","Low"))))')

first = HR + 1
for i, r in enumerate(RISKS):
    row = first + i
    (rid, title, desc, asset, cat, owner, L, I, existing, annex, treat, plan, status, due, rl, ri) = r
    values = {
        1: rid, 2: title, 3: desc, 4: asset, 5: cat, 6: owner, 7: L, 8: I,
        9: f"=G{row}*H{row}", 10: rating(f"I{row}"),
        11: existing, 12: annex, 13: treat, 14: plan, 15: status, 16: due,
        17: f'=IF(OR(O{row}="Closed",O{row}="Accepted",P{row}=""),0,MAX(0,TODAY()-P{row}))',
        18: f'=IF(Q{row}>0,"Yes","No")',
        19: rl, 20: ri, 21: f"=S{row}*T{row}", 22: rating(f"U{row}"),
        23: D(2026, 9, 30),
    }
    for j, v in values.items():
        c = ws.cell(row=row, column=j, value=v)
        c.font = f()
        c.border = BORDER
        c.alignment = CENTER if j in (1, 7, 8, 9, 10, 13, 15, 16, 17, 18, 19, 20, 21, 22, 23) else WRAP
        if cols[j - 1][2] == "in":
            c.fill = INPUT_FILL
        if j in (16, 23):
            c.number_format = "dd.mm.yyyy"
last = first + len(RISKS) - 1

ws.freeze_panes = ws.cell(row=first, column=3)
ws.auto_filter.ref = f"A{HR}:{get_column_letter(len(cols))}{last}"

dv_score = DataValidation(type="list", formula1="=Lists!$A$2:$A$6", allow_blank=True)
dv_treat = DataValidation(type="list", formula1="=Lists!$B$2:$B$5", allow_blank=True)
dv_status = DataValidation(type="list", formula1="=Lists!$C$2:$C$6", allow_blank=True)
for dv in (dv_score, dv_treat, dv_status):
    ws.add_data_validation(dv)
rng_end = 100
for col in "GHST":
    dv_score.add(f"{col}{first}:{col}{rng_end}")
dv_treat.add(f"M{first}:M{rng_end}")
dv_status.add(f"O{first}:O{rng_end}")

COLORS = {"Critical": ("C00000", "FFFFFF"), "High": ("F4B183", "000000"),
          "Medium": ("FFE699", "000000"), "Low": ("C6EFCE", "000000")}
for col in ("J", "V"):
    for name, (bg, fg) in COLORS.items():
        ws.conditional_formatting.add(f"{col}{first}:{col}{rng_end}",
            CellIsRule(operator="equal", formula=[f'"{name}"'],
                       fill=PatternFill("solid", fgColor=bg), font=Font(name=FONT, bold=True, color=fg)))
ws.conditional_formatting.add(f"R{first}:R{rng_end}",
    CellIsRule(operator="equal", formula=['"Yes"'], fill=PatternFill("solid", fgColor="F8CBAD"),
               font=Font(name=FONT, bold=True, color="C00000")))

# ---------------- Summary ----------------
su = wb.create_sheet("Summary", 2)
su["A1"] = "Risk Summary"; su["A1"].font = f(bold=True, size=14, color="1F3864")
su["A2"] = "Calculated from the Risk Register (rows 5 to 100). Portfolio exercise."; su["A2"].font = f(italic=True)
R = "'Risk Register'!"
rr = lambda c: f"{R}${c}$5:${c}$100"
kpis = [
 ("Total risks", f'=COUNTA({rr("A")})'),
 ("Open risks (not Closed or Accepted)", f'=COUNTA({rr("A")})-COUNTIF({rr("O")},"Closed")-COUNTIF({rr("O")},"Accepted")'),
 ("Overdue remediations", f'=COUNTIF({rr("R")},"Yes")'),
 ("Critical or High (inherent)", f'=COUNTIF({rr("J")},"Critical")+COUNTIF({rr("J")},"High")'),
 ("Critical or High (residual target)", f'=COUNTIF({rr("V")},"Critical")+COUNTIF({rr("V")},"High")'),
 ("Average inherent score", f'=IFERROR(AVERAGE({rr("I")}),0)'),
 ("Average residual score", f'=IFERROR(AVERAGE({rr("U")}),0)'),
]
su["A4"] = "KPI"; su["B4"] = "Value"
for c in (su["A4"], su["B4"]):
    c.font = f(bold=True, color="FFFFFF"); c.fill = HEAD_FILL; c.border = BORDER
for i, (k, formula) in enumerate(kpis, 5):
    su.cell(row=i, column=1, value=k).font = f()
    c = su.cell(row=i, column=2, value=formula); c.font = f(bold=True); c.alignment = Alignment(horizontal="center")
    if "Average" in k:
        c.number_format = "0.0"
    for j in (1, 2):
        su.cell(row=i, column=j).border = BORDER

su["D4"] = "Status"; su["E4"] = "Count"
for c in (su["D4"], su["E4"]):
    c.font = f(bold=True, color="FFFFFF"); c.fill = HEAD_FILL; c.border = BORDER
for i, st in enumerate(lists["C"][1], 5):
    su.cell(row=i, column=4, value=st).font = f()
    c = su.cell(row=i, column=5, value=f'=COUNTIF({rr("O")},D{i})'); c.font = f(); c.alignment = Alignment(horizontal="center")
    for j in (4, 5):
        su.cell(row=i, column=j).border = BORDER

su["G4"] = "Rating"; su["H4"] = "Inherent"; su["I4"] = "Residual"
for c in (su["G4"], su["H4"], su["I4"]):
    c.font = f(bold=True, color="FFFFFF"); c.fill = HEAD_FILL; c.border = BORDER
for i, name in enumerate(["Critical", "High", "Medium", "Low"], 5):
    su.cell(row=i, column=7, value=name).font = f(bold=True)
    su.cell(row=i, column=7).fill = PatternFill("solid", fgColor=COLORS[name][0])
    su.cell(row=i, column=7).font = f(bold=True, color=COLORS[name][1])
    su.cell(row=i, column=8, value=f'=COUNTIF({rr("J")},G{i})')
    su.cell(row=i, column=9, value=f'=COUNTIF({rr("V")},G{i})')
    for j in (7, 8, 9):
        su.cell(row=i, column=j).border = BORDER
        if j > 7:
            su.cell(row=i, column=j).font = f(); su.cell(row=i, column=j).alignment = Alignment(horizontal="center")

# Heatmap: rows = likelihood 5..1, columns = impact 1..5
top = 15
su.cell(row=top - 1, column=1, value="Inherent risk heatmap (number of risks per cell)").font = f(bold=True)
su.cell(row=top, column=1, value="Likelihood ↓ / Impact →").font = f(italic=True)
for j, imp in enumerate(range(1, 6), 2):
    c = su.cell(row=top, column=j, value=imp); c.font = f(bold=True); c.alignment = Alignment(horizontal="center"); c.border = BORDER
for i, lik in enumerate(range(5, 0, -1), top + 1):
    c = su.cell(row=i, column=1, value=lik); c.font = f(bold=True); c.alignment = Alignment(horizontal="center"); c.border = BORDER
    for j, imp in enumerate(range(1, 6), 2):
        score = lik * imp
        band = "Critical" if score >= 15 else "High" if score >= 10 else "Medium" if score >= 5 else "Low"
        c = su.cell(row=i, column=j, value=f'=COUNTIFS({rr("G")},$A{i},{rr("H")},{get_column_letter(j)}${top})')
        c.fill = PatternFill("solid", fgColor=COLORS[band][0])
        c.font = f(bold=True, size=12, color=COLORS[band][1])
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
    su.row_dimensions[i].height = 28
su.cell(row=top + 7, column=1, value="Cell colours show the default rating bands (Low 1-4, Medium 5-9, High 10-14, Critical 15-25).").font = f(italic=True, color="595959")
su.column_dimensions["A"].width = 36
for col in "BCDEFGHI":
    su.column_dimensions[col].width = 12
su.column_dimensions["D"].width = 14

# ---------------- Sources ----------------
# (section, source, url, what it supports, used in, access)
SOURCES = [
 ("Organisation", "Saarland University: My new working environment (Saarbrücken and Homburg campuses)",
  "https://uni-saarland.de/en/page/newly-appointed-professors/your-professorship/working-environment.html",
  "Campuses in Saarbrücken and Homburg; Faculty of Medicine at the Medical Center in Homburg; University IT Centre (HIZ) with HR and Finance at the Meerwiesertalweg site; IT Service Desk in the Campus Center",
  "Read Me; R03, R07, R09", "Free web page"),
 ("Organisation", "Saarland University: Welche weiteren Anlaufstellen gibt es an der UdS?",
  "https://www.uni-saarland.de/page/neuberufen/persoenliches/anlaufstellen.html",
  "HIZ is the shared IT service provider of htw saar and Saarland University",
  "Read Me; owner role 'Head of University IT Centre (HIZ)'", "Free web page"),
 ("Organisation", "Hochschul-IT-Zentrum Saar (HIZ) website", "https://www.hiz-saarland.de/start/",
  "Official site of the IT centre (listed for reference; its content was not used)", "Owner role", "Free web page"),
 ("Organisation", "Saarland University: The Homburg campus",
  "https://uni-saarland.de/en/studies/saarland/the-homburg-campus.html",
  "Medical campus in Homburg", "R03", "Free web page"),
 ("Organisation", "Student council computer science: Internet @ Uni",
  "https://cs.fs.uni-saarland.de/en/special/internet/",
  "University offers eduroam WiFi and a VPN for remote access", "R02, R13", "Free web page"),
 ("Organisation", "Wikipedia: Saarland University", "https://en.wikipedia.org/wiki/Saarland_University",
  "Background only: founded 1948, six faculties, about 16,300 students", "Background, not in the register", "Free web page"),
 ("Organisation", "DFN: The national research and education network", "https://dfn.de/en/network",
  "German research network (X-WiN). A Saarland University uplink via DFN is assumed, not confirmed by a university page", "R05 (assumption)", "Free web page"),
 ("Standard", "ISO/IEC 27001:2022 Information security management systems: Requirements (with Amd 1:2024)",
  "https://www.iso.org/standard/27001",
  "Annex A with 93 controls; control numbers and names in column L", "Column L; Statement of Applicability later", "Paid standard (about CHF 155), no copy included"),
 ("Standard", "ISO/IEC 27002:2022 Information security controls", "https://www.iso.org/standard/75652.html",
  "Guidance for each Annex A control", "Column L", "Paid standard, no copy included"),
 ("Method", "BSI-Standards overview (incl. BSI-Standard 200-3 risk management)",
  "https://www.bsi.bund.de/DE/Themen/Unternehmen-und-Organisationen/Standards-und-Zertifizierung/IT-Grundschutz/BSI-Standards/bsi-standards_node.html",
  "German government method for risk analysis, used as background. The 5x5 scales here are this exercise's own, not BSI's",
  "Scoring Method", "Free download from BSI"),
 ("Method", "BSI-Standard 200-3 English PDF",
  "https://www.bsi.bund.de/SharedDocs/Downloads/EN/BSI/Grundschutz/International/bsi-standard-2003_en_pdf.html",
  "Same, English version", "Scoring Method", "Free download from BSI"),
 ("Law", "EU General Data Protection Regulation (GDPR), Regulation (EU) 2016/679",
  "https://eur-lex.europa.eu/eli/reg/2016/679/oj",
  "Art. 32 security of processing; Art. 33 breach notification, used in the impact scale ('reportable GDPR breach')",
  "Scoring Method; R03, R10, R18", "Free"),
 ("Threat context", "zm-online: Justus-Liebig-Universität ist wieder online (7 Jan 2020)",
  "https://www.zm-online.de/news/detail/justus-liebig-universitaet-ist-wieder-online",
  "Malware attack (likely Emotet) shut down Uni Gießen in December 2019", "R01", "Free web page"),
 ("Threat context", "The Record: Vice Society claims ransomware attack on HAW Hamburg (6 Mar 2023)",
  "https://therecord.media/germany-ransomware-haw-hamburg-vice-society",
  "Ransomware at HAW Hamburg (Dec 2022) and University of Duisburg-Essen (Nov 2022)", "R01", "Free web page"),
]
so = wb.create_sheet("Sources")
so["A1"] = "Sources"; so["A1"].font = f(bold=True, size=14, color="1F3864")
so["A2"] = "Checked 8 October 2026. Everything not listed here (risks, scores, owners, existing controls, dates) is invented for the exercise."
so["A2"].font = f(italic=True)
shdr = ["Type", "Source", "Link", "What it supports", "Used in", "Access"]
for j, h in enumerate(shdr, 1):
    c = so.cell(row=4, column=j, value=h); c.font = f(bold=True, color="FFFFFF"); c.fill = HEAD_FILL; c.border = BORDER
for i, src in enumerate(SOURCES, 5):
    for j, v in enumerate(src, 1):
        c = so.cell(row=i, column=j, value=v); c.font = f(); c.alignment = WRAP; c.border = BORDER
        if j == 3:
            c.hyperlink = v
            c.font = f(color="0563C1", underline="single")
for col, w in zip("ABCDEF", (14, 40, 45, 55, 24, 22)):
    so.column_dimensions[col].width = w

if len(sys.argv) > 2:
    with open(sys.argv[2], "w") as md:
        md.write("# Sources for the Saarland University risk register\n\n")
        md.write("Portfolio exercise, not an official university document. Sources checked on 8 October 2026. "
                 "Only the facts below come from public sources; all risks, scores, owners, existing controls, "
                 "statuses and dates in the register are invented.\n\n")
        for sec in dict.fromkeys(s[0] for s in SOURCES):
            md.write(f"## {sec}\n\n")
            for _, name, url, what, used, access in (s for s in SOURCES if s[0] == sec):
                md.write(f"- [{name}]({url})  \n  Supports: {what}.  \n  Used in: {used}. Access: {access}.\n")
            md.write("\n")
        md.write("## Files in this folder\n\n"
                 "- `saarland-university-risk-register.xlsx`: the register (sheets Read Me, Risk Register, Summary, Scoring Method, Lists, Sources).\n"
                 "- `build_register.py`: Python script that generates the workbook and this file (`python3 build_register.py out.xlsx sources.md`).\n"
                 "- `sources.md`: this list.\n\n"
                 "The ISO standards are paid and cannot be shared. The BSI and GDPR documents are free to download from the links above.\n")

wb.save(OUT)
print("saved", OUT, len(RISKS), "risks")
